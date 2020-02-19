from fpdf import FPDF
from docx import Document


def gen():
    with open('tekst.txt', 'r') as file:
        content = []
        for line in file:
            content.append(line)
#        content = file.read()
        return content
#        print(content)


def pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', size = 10)
    for line in gen():
        pdf.cell(500, 10, txt = line, ln = 1, align = 'L')
    pdf.output('something.pdf')

def word():
    doc = Document()
    for line in gen():
        doc.add_paragraph(line)
    doc.save('something.doc')

if __name__ == '__main__':
    gen()
    doc_form = input('Choose your doc format (pdf/word): ').upper()
    if doc_form == 'PDF':
        pdf()
        print('Converting txt to .pdf file')
    elif doc_form == 'WORD':
        word()
        print('Converting txt to .doc file')
    else:
        print('Do not support this format yet')
        quit()






'''
class DocGenInterface():
    def gen():
        with open('tekst.txt', r) as file:
            content = file.read()
            print(content)


class pdf(DocGenInterface):
    def gen(self):
        print(f'generating .pdf file')

class word(DocGenInterface):
    def gen(self):
        print('generating .doc file')

class DocGenFacotry():
    @staticmethod
    def get_gen(gen_str):
        method = {
            'word': word(),
            'pdf': pdf()
        }
        gen = method.get(gen_str)
        if gen == None:
            assert 0, f'Cannot convert to <{gen_str}> format'
        else:
            return gen


if __name__ == '__main__':
    g = DocGenFacotry()
    option = input('Enter document format: ')
    gen_option = g.get_gen(option)
    gen_option.gen()

'''